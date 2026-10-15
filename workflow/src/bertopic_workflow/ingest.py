"""Content-sniffed table readers and provenance-preserving database adapters."""
from __future__ import annotations

from collections import Counter
import csv
import hashlib
import html
import io
from pathlib import Path
import re
import unicodedata
import zipfile

from .common import WorkflowError, sha256, write_json, write_table
from .config import FIELDS, InputFile

ALIASES = {
    "title": ["title", "article title", "document title", "Title-题名", "题名", "标题", "篇名", "文献题名", "TI"],
    "abstract": ["abstract", "summary", "Summary-摘要", "摘要", "中文摘要", "内容摘要", "AB"],
    "year": ["year", "publication year", "Year-年", "年", "年份", "出版年", "发表年份", "PY"],
    "date": ["publication date", "PubTime-发表时间", "发表时间", "出版日期", "日期", "DP"],
    "doi": ["doi", "DOI-DOI", "digital object identifier", "DI"],
    "journal": ["journal", "journal/book", "source title", "Source-文献来源", "刊名", "期刊", "来源", "文献来源", "来源刊名", "SO", "JT", "JO", "JF"],
    "authors": ["authors", "author", "Author-作者", "作者", "第一作者", "AU", "FAU"],
    "keywords": ["keywords", "author keywords", "Keyword-关键词", "关键词", "关键字", "DE", "KW"],
    "institutions": ["affiliations", "institutions", "institution", "Organ-机构", "Organ-单位", "机构", "作者单位", "单位", "组织", "C3", "AD"],
    "addresses": ["addresses", "address", "C1"],
    "countries": ["country", "countries", "country/region", "国家", "国家地区", "国家/地区"],
    "category": ["wos categories", "subject", "category", "CLC-中图分类号", "中图分类号", "分类号", "学科", "学科分类", "WC"],
    "research_areas": ["research areas", "研究领域", "SC"],
    "document_type": ["document type", "SrcDatabase-来源库", "文献类型", "类型", "DT", "TY"],
    "publication_type": ["publication type", "publication_types", "PT"],
    "pmid": ["pmid", "pubmed id"],
    "wos_id": ["UT (Unique WOS ID)", "accession number", "wos id", "UT"],
}


def key(value):
    return re.sub(r"[^\w]", "", unicodedata.normalize("NFKC", str(value)).casefold().lstrip("\ufeff"))


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def clean_text(value):
    text = html.unescape(clean(value))
    text = re.sub(r"<正>|</?[A-Za-z][^>]*>", " ", text)
    text = re.sub(r"\s+(?:Copyright\s+(?:©|\(C\)|\d{4})|©\s*\d{4}|\(C\)\s*\d{4}).*$", "", text, flags=re.I)
    return clean(text)


def doi(value):
    value = re.sub(r"^(?:https?://(?:dx\.)?doi\.org/|doi\s*:\s*)", "", clean(value).casefold())
    return value.rstrip(".") if re.fullmatch(r"10\.\d{4,9}/\S+", value) else ""


def decode_file(path, encoding=None):
    data = Path(path).read_bytes()
    encodings = [encoding] if encoding else (["utf-16"] if data[:2] in [b"\xff\xfe", b"\xfe\xff"] else ["utf-8-sig", "gb18030", "cp1252"])
    for codec in encodings:
        try:
            return data.decode(codec), codec
        except (UnicodeDecodeError, LookupError):
            continue
    raise WorkflowError("ENCODING", "无法解码输入，请显式指定 encoding。", details={"file": Path(path).name})


def mapping_for(headers, overrides):
    normalized = [key(h) for h in headers]
    result = {}
    for field, aliases in ALIASES.items():
        wanted = {key(a) for a in aliases}
        matches = [i for i, name in enumerate(normalized) if name in wanted]
        if field in overrides:
            matches = [i for i, name in enumerate(headers) if clean(name) == overrides[field]]
            if len(matches) != 1:
                raise WorkflowError("MAPPING", "显式映射列不存在或不唯一。", details={"field": field})
        if len(matches) > 1:
            raise WorkflowError("MAPPING_AMBIGUOUS", "多个列对应同一字段，请显式选择。", details={"field": field, "columns": [headers[i] for i in matches]})
        if matches:
            result[field] = matches[0]
    return result


def identify(headers, filename):
    h = {key(v) for v in headers}
    if "title题名" in h or "srcdatabase来源库" in h:
        return "CNKI"
    if "utuniquewosid" in h or {"articletitle", "sourcetitle"} <= h:
        return "WoS"
    if "pmid" in h or "pubmedid" in h:
        return "PubMed"
    name = filename.casefold()
    for source, pattern in [("VIP", r"cqvip|维普|\bvip\b"), ("WanFang", r"wanfang|万方"),
                            ("CNKI", r"cnki|知网"), ("WoS", r"wos|web.of.science"), ("PubMed", r"pubmed")]:
        if re.search(pattern, name):
            return source
    return "Generic"


def choose_header(rows, options):
    candidates = []
    indices = [options.header_row - 1] if options.header_row else range(min(30, len(rows)))
    for index in indices:
        if index >= len(rows):
            raise WorkflowError("HEADER", "header_row 超出表格范围。")
        headers = [clean(v) for v in rows[index]]
        try:
            mapping = mapping_for(headers, options.mapping)
        except WorkflowError:
            if options.header_row or any(key(h) in {key(a) for a in ALIASES['title']} for h in headers):
                raise
            continue
        if "title" in mapping or "abstract" in mapping:
            score = 6 * ("title" in mapping) + 5 * ("abstract" in mapping) + len(mapping)
            candidates.append((score, -index, headers, mapping))
    if not candidates:
        raise WorkflowError("HEADER", "未识别题名或摘要列；请使用 inspect 并配置 header_row/mapping。")
    _, neg_index, headers, mapping = max(candidates, key=lambda c: c[:2])
    if len([h for h in headers if h]) != len(set(h for h in headers if h)):
        raise WorkflowError("HEADER_AMBIGUOUS", "存在重复表头，请提供规范表格。")
    return -neg_index, headers, mapping


def html_tables(text):
    from lxml import html as lhtml
    parser = lhtml.HTMLParser(no_network=True)
    document = lhtml.fromstring(text, parser=parser)
    tables = []
    for number, table in enumerate(document.xpath('//table')):
        grid, pending = [], {}
        for row_index, row in enumerate(table.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')):
            cells = {}
            for col, (until, value) in list(pending.items()):
                if until > row_index:
                    cells[col] = value
                else:
                    pending.pop(col)
            col = 0
            for cell in row.xpath('./th|./td'):
                while col in cells:
                    col += 1
                value = clean(cell.text_content())
                try:
                    width, height = int(cell.get('colspan', '1')), int(cell.get('rowspan', '1'))
                except ValueError:
                    raise WorkflowError("TABLE_FORMAT", "HTML 表格的合并单元格无效。") from None
                if not 1 <= width <= 256 or not 1 <= height <= 10000:
                    raise WorkflowError("TABLE_LIMIT", "HTML 合并单元格范围过大。")
                for offset in range(width):
                    cells[col + offset] = value
                    if height > 1:
                        pending[col + offset] = (row_index + height, value)
                col += width
            grid.append([cells.get(i, '') for i in range(max(cells, default=-1) + 1)])
        tables.append((f"table:{number + 1}", grid))
    return tables


def medline(text):
    records, record, tag = [], {}, None
    for line in text.splitlines():
        m = re.match(r"^([A-Z0-9]{2,4})\s*-\s?(.*)$", line)
        if m:
            tag, value = m.groups()
            if tag == "PMID" and record:
                records.append(record)
                record = {}
            record.setdefault(tag, []).append(value)
        elif line.startswith(' ') and tag:
            record[tag][-1] += ' ' + line.strip()
    if record:
        records.append(record)
    output = []
    for item in records:
        ids = item.get('AID', []) + item.get('LID', [])
        output.append({"title": ' '.join(item.get('TI', [])), "abstract": ' '.join(item.get('AB', [])),
                       "pmid": '; '.join(item.get('PMID', [])), "year": '; '.join(item.get('DP', [])),
                       "doi": next((re.sub(r'\s*\[doi\]$', '', v) for v in ids if '[doi]' in v), ''),
                       "journal": '; '.join(item.get('JT', item.get('TA', []))),
                       "authors": '; '.join(item.get('FAU', item.get('AU', []))),
                       "institutions": '; '.join(item.get('AD', [])),
                       "document_type": '; '.join(item.get('PT', [])),
                       "keywords": '; '.join(item.get('OT', []))})
    return output


def ris(text):
    records, record, tag = [], {}, None
    for line in text.splitlines():
        m = re.match(r'^([A-Z0-9]{2})\s+-\s?(.*)$', line)
        if m:
            tag, value = m.groups()
            if tag == 'TY' and record:
                records.append(record)
                record = {}
            if tag == 'ER':
                records.append(record)
                record, tag = {}, None
            else:
                record.setdefault(tag, []).append(value)
        elif line.strip() and tag:
            record[tag][-1] += ' ' + line.strip()
    if record:
        records.append(record)
    specs = {'title': ['TI', 'T1'], 'abstract': ['AB', 'N2'], 'year': ['PY', 'Y1'],
             'doi': ['DO'], 'journal': ['JO', 'JF', 'T2'], 'authors': ['AU', 'A1'],
             'keywords': ['KW'], 'document_type': ['TY']}
    return [{field: '; '.join(next((item[tag] for tag in tags if tag in item), []))
             for field, tags in specs.items()} for item in records]


def xml_pubmed(text):
    from defusedxml.ElementTree import fromstring
    root = fromstring(text)
    output = []
    for article in root.findall('.//PubmedArticle'):
        def value(xpath, article=article):
            node = article.find(xpath)
            return clean(''.join(node.itertext())) if node is not None else ''
        output.append({'title': value('.//ArticleTitle'),
                       'abstract': ' '.join(clean(''.join(n.itertext())) for n in article.findall('.//AbstractText')),
                       'pmid': value('.//PMID'), 'journal': value('.//Journal/Title'),
                       'year': value('.//PubDate/Year') or value('.//PubDate/MedlineDate'),
                       'doi': next((clean(n.text) for n in article.findall('.//ArticleId') if n.get('IdType') == 'doi'), ''),
                       'authors': '; '.join(clean(' '.join(filter(None, [n.findtext('LastName'), n.findtext('ForeName')]))) for n in article.findall('.//Author')),
                       'institutions': '; '.join(clean(''.join(n.itertext())) for n in article.findall('.//Affiliation'))})
    return output


def read_input(options: InputFile, limits):
    path = Path(options.path)
    with path.open('rb') as stream:
        prefix = stream.read(1024)
    fmt, encoding, tables = None, None, []
    structured = None
    if prefix.startswith(b'PK\x03\x04'):
        fmt = 'xlsx'
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > limits.max_uncompressed_mb * 1024**2:
                raise WorkflowError('INPUT_LIMIT', 'Excel 解压大小超出限制。')
        import openpyxl
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                if options.sheet is not None and options.sheet not in [sheet.title, workbook.worksheets.index(sheet)]:
                    continue
                if sheet.max_row and sheet.max_row > limits.max_records + 30:
                    raise WorkflowError('INPUT_LIMIT', 'Excel 工作表行数超出限制。')
                tables.append((sheet.title, list(sheet.iter_rows(values_only=True))))
        finally:
            workbook.close()
    elif prefix.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
        fmt = 'xls'
        import xlrd
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            for index, sheet in enumerate(workbook.sheets()):
                if options.sheet is not None and options.sheet not in [sheet.name, index]:
                    continue
                if sheet.nrows > limits.max_records + 30:
                    raise WorkflowError('INPUT_LIMIT', 'Excel 工作表行数超出限制。')
                tables.append((sheet.name, [sheet.row_values(i) for i in range(sheet.nrows)]))
        finally:
            workbook.release_resources()
    else:
        text, encoding = decode_file(path, options.encoding)
        if re.search(r'<(?:html|table)\b', text[:4096], re.I):
            fmt, tables = 'html_table', html_tables(text)
        elif re.search(r'^PMID\s*-', text, re.M):
            fmt, structured = 'medline', medline(text)
        elif re.search(r'^TY\s+-', text, re.M):
            fmt, structured = 'ris', ris(text)
        elif '<PubmedArticle' in text[:4096]:
            fmt, structured = 'pubmed_xml', xml_pubmed(text)
        else:
            fmt = 'delimited'
            delimiter = options.delimiter
            if delimiter is None:
                try:
                    delimiter = csv.Sniffer().sniff(text[:65536], delimiters=',\t;').delimiter
                except csv.Error:
                    delimiter = '\t' if '\t' in text.splitlines()[0] else ','
            csv.field_size_limit(16 * 1024**2)
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
            if len(rows) > limits.max_records + 30:
                raise WorkflowError('INPUT_LIMIT', '表格行数超出限制。')
            tables = [('table:1', rows)]
    if structured is not None:
        if not structured:
            raise WorkflowError('INPUT_EMPTY', '文献格式中没有可读取记录。')
        headers = list(structured[0])
        tables = [('records', [headers] + [[r.get(h, '') for h in headers] for r in structured])]
    if fmt == 'html_table' and options.sheet is not None:
        tables = [(name, rows) for i, (name, rows) in enumerate(tables) if options.sheet in {name, i}]
    candidates, errors = [], []
    for table_name, rows in tables:
        try:
            header, headers, mapping = choose_header(rows, options)
            candidates.append((len(mapping), table_name, rows, header, headers, mapping))
        except WorkflowError as error:
            errors.append(error)
    if not candidates:
        if errors:
            raise errors[0]
        raise WorkflowError('SHEET', '指定工作表不存在或没有表格。')
    if len(candidates) > 1:
        raise WorkflowError('SHEET_AMBIGUOUS', '多个表格含文献列，请明确指定 sheet。', details={'sheets': [c[1] for c in candidates]})
    _, table_name, rows, header, headers, mapping = candidates[0]
    source = options.source if options.source != 'auto' else ('PubMed' if fmt in {'medline', 'pubmed_xml'} else identify(headers, path.name))
    records, blank_rows = [], 0
    for index, row in enumerate(rows[header + 1:], header + 2):
        if not any(clean(v) for v in row):
            blank_rows += 1
            continue
        if len(row) > len(headers) and any(clean(v) for v in row[len(headers):]):
            raise WorkflowError('ROW_WIDTH', '表格存在超出表头的非空列。', details={'row': index})
        record = {field: clean(row[col]) if col < len(row) else '' for field, col in mapping.items()}
        record['_row'] = index
        records.append(record)
    return records, {'source': source, 'format': fmt, 'encoding': encoding, 'sheet': table_name,
                     'header_row': header + 1, 'headers': headers,
                     'mapping': {field: headers[index] for field, index in mapping.items()},
                     'records': len(records), 'blank_rows': blank_rows,
                     'warnings': ['未提供摘要，使用题名建模。'] if 'abstract' not in mapping else []}


def prepare(config, destination, check_cancel=lambda: None, progress=lambda **kwargs: None):
    records, rejected, duplicates, manifests, ids = [], [], [], [], set()
    dedup = {}
    total_read = 0
    for input_index, options in enumerate(config.inputs):
        check_cancel()
        file = Path(options.path)
        before = sha256(file)
        raw_rows, info = read_input(options, config.import_options)
        after = sha256(file)
        if before != after:
            raise WorkflowError('INPUT_CHANGED', '输入在读取过程中发生改变。')
        info.update(path=str(file), sha256=before, input_index=input_index)
        manifests.append(info)
        total_read += len(raw_rows)
        if total_read > config.import_options.max_records:
            raise WorkflowError('INPUT_LIMIT', '联合语料超出记录数上限。')
        for raw in raw_rows:
            record_id = hashlib.sha256(f'{before}\0{input_index}\0{info["sheet"]}\0{raw["_row"]}'.encode()).hexdigest()[:24]
            if record_id in ids:
                raise WorkflowError('RECORD_ID', '文献 ID 出现碰撞。')
            ids.add(record_id)
            title, abstract = clean_text(raw.get('title')), clean_text(raw.get('abstract'))
            if not title and not abstract:
                reject = {'record_id': record_id, 'source': info['source'], 'source_file': file.name,
                          'source_row': raw['_row'], 'reason': 'missing_title_and_abstract', 'raw': raw}
                if config.import_options.invalid_rows == 'error':
                    raise WorkflowError('DOCUMENT_EMPTY', '文献题名与摘要均为空。', details={'file': file.name, 'row': raw['_row']})
                rejected.append(reject)
                continue
            separator = (' ' if title.endswith(('.', '!', '?', '。', '！', '？')) else '. ') if title and abstract else ''
            text = (title + separator + abstract).strip()
            cjk, latin = len(re.findall(r'[\u3400-\u9fff]', text)), len(re.findall('[A-Za-z]', text))
            language = 'zh' if cjk / max(1, cjk + latin) > .2 else ('en' if latin else 'other')
            year_raw = raw.get('year') or raw.get('date', '')
            year_match = re.search(r'(?<!\d)(?:1[5-9]|20|21)\d{2}(?!\d)', year_raw)
            year = int(year_match[0]) if year_match else None
            item = {field: raw.get(field, '') for field in sorted(FIELDS) if field not in {'title', 'abstract', 'year'}}
            item.update(record_id=record_id, source=info['source'], source_file=file.name,
                        source_sha256=before, source_sheet=info['sheet'], source_row=raw['_row'],
                        title=title, abstract=abstract, text=text, has_abstract=bool(abstract),
                        year=year, year_raw=year_raw, language_of_text=language, doi=doi(raw.get('doi')),
                        text_sha256=hashlib.sha256(text.encode()).hexdigest())
            item['document_type'] = item['document_type'] or item['publication_type']
            item['doi_raw'] = raw.get('doi', '')
            item['metadata_origin'] = {'countries': 'exported_field' if item['countries'] else 'missing',
                                       'institutions': 'exported_field' if item['institutions'] else 'missing'}
            if item['addresses'] and (not item['countries'] or not item['institutions']):
                from .metadata import address_groups
                countries, institutions = address_groups(item['addresses'])
                for field, values in [('countries', countries), ('institutions', institutions)]:
                    if not item[field] and values:
                        item[field] = '; '.join(values)
                        item['metadata_origin'][field] = 'parsed_address'
            policy = config.import_options.deduplication
            dedup_key = item['text_sha256'] if policy == 'exact_text' else item['doi'] if policy == 'doi' else ''
            scope = item['source'] if config.import_options.deduplication_scope == 'within_source' else '*'
            if dedup_key and (scope, dedup_key) in dedup:
                duplicates.append({'record_id': record_id, 'kept_record_id': dedup[(scope, dedup_key)],
                                   'rule': policy, 'scope': scope, 'source_file': file.name, 'source_row': raw['_row']})
                continue
            if dedup_key:
                dedup[(scope, dedup_key)] = record_id
            records.append(item)
        progress(file_index=input_index + 1, files=len(config.inputs), records=len(records))
    if len(records) < config.import_options.min_records:
        raise WorkflowError('CORPUS_TOO_SMALL', '有效记录数不足；请检查字段映射及导出内容。', details={'records': len(records)})
    counts = Counter(r['text_sha256'] for r in records)
    for record in records:
        record['exact_text_record_count'] = counts[record['text_sha256']]
    manifest = {'input_files': manifests, 'read_records': total_read, 'retained_records': len(records),
                'quarantined_records': len(rejected), 'deduplicated_records': len(duplicates),
                'source_counts': dict(Counter(r['source'] for r in records)),
                'missing_abstracts': sum(not r['has_abstract'] for r in records),
                'missing_years': sum(r['year'] is None for r in records),
                'duplicate_text_records': sum(n - 1 for n in counts.values()),
                'deduplication': config.import_options.deduplication,
                'deduplication_scope': config.import_options.deduplication_scope}
    write_table(destination, 'corpus', records)
    write_table(destination, 'quarantined_rows', rejected, ['record_id', 'source', 'source_file', 'source_row', 'reason', 'raw'])
    write_table(destination, 'duplicate_lineage', duplicates, ['record_id', 'kept_record_id', 'rule', 'scope', 'source_file', 'source_row'])
    write_json(Path(destination) / 'input_manifest.json', manifest)
    return manifest
